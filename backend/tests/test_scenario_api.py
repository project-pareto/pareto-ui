"""Acceptance checks use isolated storage and the real model/report pipeline."""
from copy import deepcopy
import asyncio
import importlib
import json
import os
import shutil
from pathlib import Path
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from fastapi import BackgroundTasks, FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pyomo.environ import SolverFactory
from tinydb import Query
from scenario_fixtures import simple_scenario, map_files, legacy_units_sheet
from app.internal.scenarios.inputs import read_inputs, write_inputs
from app.internal.workbooks.reader import get_data


class ScenarioApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory()
        cwd = os.getcwd()
        with patch.dict(os.environ, {'PARETO_DATA_BASEDIR': cls.directory.name, 'PARETO_LOG_DIR': cls.directory.name}):
            cls.module = importlib.import_module('app.internal.scenario_handler')
            cls.routes = importlib.import_module('app.routers.scenarios')
            cls.runner = importlib.import_module('app.internal.optimization.strategic_model')
        os.chdir(cwd)

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        cwd = os.getcwd()
        self.handler = self.module.ScenarioHandler(data_basedir=self.temp.name, log_dir=self.temp.name)
        os.chdir(cwd)
        self.addCleanup(self.handler._db.close)
        for module in (self.routes, self.runner):
            replacement = patch.object(module, 'scenario_handler', self.handler)
            replacement.start(); self.addCleanup(replacement.stop)
        app = FastAPI(); app.include_router(self.routes.router)
        self.client = TestClient(app)
        self.example, _ = simple_scenario(self.temp.name)
        write_inputs(self.example['data_input'], self.handler.get_excelsheet_path(1))
        self.example = self.handler.update_scenario(self.example)

    def test_validation_is_invalidated_and_stale_or_running_edits_are_rejected(self):
        revision = self.example['input_revision']
        checked = self.client.get('/validate_scenario/1')
        self.assertEqual(checked.json()['model_check'], 'passed', checked.text)
        changed = self.client.post('/planning_horizon/1', json={'periods': ['T01', 'T02', 'T03'], 'revision': revision})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertEqual(changed.json()['validation']['state'], 'needs_input')
        self.assertEqual(changed.json()['validation']['model_check'], 'not_run')
        self.assertEqual(changed.json()['validation']['feasibility'], 'not_run')
        self.assertEqual(self.client.post('/advance_to_optimization_setup/1').status_code, 409)
        self.assertEqual(self.client.post('/planning_horizon/1', json={'periods': ['T01'], 'revision': revision}).status_code, 409)
        self.assertEqual(self.client.post('/run_model', json={'scenario': self.example}).status_code, 409)
        current = self.handler.get_scenario(1)
        self.handler.add_background_task(1)
        self.assertEqual(self.client.post('/update_excel', json={'id': 1, 'tableKey': 'PadRates', 'updatedTable': {}, 'revision': current['input_revision']}).status_code, 409)
        self.assertEqual(self.handler.get_scenario(1)['input_revision'], current['input_revision'])

    def test_payload_models_match_current_scenario_and_readiness_responses(self):
        from app.schemas.scenario import Scenario
        from app.schemas.validation import ScenarioValidation

        response = self.client.get('/get_scenario/1')
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(Scenario.model_validate(payload).to_payload(), payload)
        response = self.client.get('/scenario_readiness/1')
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(ScenarioValidation.model_validate(payload).to_payload(), payload)

    def create_zero_id_scenario(self):
        scenario = deepcopy(self.example)
        scenario['id'] = 0
        write_inputs(scenario['data_input'], self.handler.get_excelsheet_path(0))
        return self.handler.update_scenario(scenario)

    def test_zero_id_update_rejects_stale_and_running_edits_without_writes(self):
        current = self.create_zero_id_scenario()
        workbook = Path(self.handler.get_excelsheet_path(0))
        original_bytes = workbook.read_bytes()
        stale = deepcopy(current)
        stale['input_revision'] = 'stale-revision'
        stale['name'] = 'Should not be saved'
        response = self.client.post('/update', json={'updatedScenario': stale})
        self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(self.handler.get_scenario(0), current)

        self.handler.add_background_task(0)
        running = deepcopy(current)
        running['data_input']['df_parameters']['PadRates']['T01'] = [42]
        response = self.client.post('/update', json={'updatedScenario': running, 'propagateChanges': 'json'})
        self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(self.handler.get_scenario(0), current)
        self.assertEqual(workbook.read_bytes(), original_bytes)

    def test_zero_id_table_and_map_updates_survive_reload_and_export(self):
        current = self.create_zero_id_scenario()
        current['data_input']['df_parameters']['PadRates']['T01'] = [42]
        response = self.client.post('/update', json={'updatedScenario': current})
        self.assertEqual(response.status_code, 200, response.text)
        current = response.json()['data']
        self.assertEqual(read_inputs(self.handler.get_excelsheet_path(0))['df_parameters']['PadRates']['T01'], [42])

        current['data_input']['map_data']['all_nodes']['K1']['DisposalOperationalCost'] = 2
        response = self.client.post('/update', json={'updatedScenario': current, 'propagateChanges': 'map'})
        self.assertEqual(response.status_code, 200, response.text)
        exported = read_inputs(self.handler.get_excelsheet_path(0))
        self.assertEqual(exported['df_parameters']['DisposalOperationalCost']['VALUE'], [2])
        self.assertEqual(exported['df_parameters']['PadRates']['T01'], [42])
        self.assertEqual(self.handler.get_scenario(0), response.json()['data'])

    def test_workshop_sra_diagrams_use_exact_asset_filename_case(self):
        scenario = deepcopy(self.example)
        scenario['name'] = 'Workshop SRA'
        self.handler.update_scenario(scenario)
        real_copy = shutil.copyfile

        def case_sensitive_copy(source, destination):
            source = Path(source)
            # macOS often resolves the wrong case; Linux and packaged assets do not.
            if source.name not in {path.name for path in source.parent.iterdir()}:
                raise FileNotFoundError(source)
            return real_copy(source, destination)

        for kind, filename, directory in (
            ('input', 'Workshop SRA', self.handler.input_diagrams_path),
            ('output', None, self.handler.output_diagrams_path),
        ):
            with self.subTest(kind=kind):
                target = Path(directory) / '1.png'
                target.unlink(missing_ok=True)
                with patch.object(self.module.shutil, 'copyfile', side_effect=case_sensitive_copy):
                    saved = self.handler.check_for_diagram(1, filename)
                self.assertTrue(target.is_file())
                source = Path(self.module.__file__).parent / 'assets' / f'workshop_SRA_{kind}.png'
                self.assertEqual(target.read_bytes(), source.read_bytes())
                self.assertEqual(saved[f'{kind}DiagramExtension'], 'png')

    def test_table_and_map_edits_survive_reload_and_export(self):
        table = {'ProductionPads': ['P1'], 'T01': [100], 'T02': [0]}
        saved = self.client.post('/update_excel', json={'id': 1, 'tableKey': 'PadRates', 'updatedTable': table}).json()
        saved['data_input']['map_data']['all_nodes']['K1']['DisposalOperationalCost'] = 2
        saved['data_input']['map_data']['all_nodes']['P1']['coordinates'][0] -= .001
        response = self.client.post('/update', json={'updatedScenario': saved, 'propagateChanges': 'map'})
        self.assertEqual(response.status_code, 200, response.text)
        current = response.json()['data']
        self.assertEqual(current['data_input']['df_parameters']['PadRates'], table)
        self.assertEqual(current['data_input']['df_parameters']['DisposalOperationalCost']['VALUE'], [2])
        exported = read_inputs(self.handler.get_excelsheet_path(1))
        self.assertEqual(exported['df_parameters']['PadRates'], table)
        self.assertEqual(exported['df_sets']['TimePeriods'], ['T01', 'T02'])
        self.assertEqual(exported['units'], current['data_input']['units'])

    def test_table_save_rejects_malformed_requests_before_reads_or_writes(self):
        valid = {'id': 1, 'tableKey': 'PadRates', 'updatedTable': {'T01': [0]}}
        cases = [(field, {key: value for key, value in valid.items() if key != field})
                 for field in valid]
        cases.extend((field, {**valid, field: value}) for field, values in (
            ('id', [False, 1.0, -1, '01', ' 1', None, 2**53]),
            ('tableKey', [False, '', '  ', None]),
            ('revision', [False, 0, [], {}]),
            ('updatedTable', [None, [], 'table', {'T01': 100}, {'T01': [True]},
                              {'T01': [[0]]}, {'T01': [{'value': 0}]}, {'T01': None}]),
        ) for value in values)
        cases.extend(('body', body) for body in ([], 'table', None))
        workbook = Path(self.handler.get_excelsheet_path(1))
        original_workbook = workbook.read_bytes()
        original_database = self.handler.scenarios_path.read_bytes()
        with patch.object(self.handler, 'get_scenario') as read, patch.object(self.handler, 'update_excel') as save:
            for field, payload in cases:
                with self.subTest(field=field, payload=payload):
                    response = self.client.post('/update_excel', json=payload)
                    self.assertEqual(response.status_code, 422, response.text)
                    errors = response.json()['detail']
                    self.assertTrue(any(error['loc'][:1] == ['body'] and
                                        (field == 'body' or field in error['loc']) for error in errors), errors)
            response = self.client.post('/update_excel', content='{"id":', headers={'Content-Type': 'application/json'})
            self.assertEqual(response.status_code, 422, response.text)
            self.assertEqual(response.json()['detail'][0]['type'], 'json_invalid')
            read.assert_not_called()
            save.assert_not_called()
        self.assertEqual(workbook.read_bytes(), original_workbook)
        self.assertEqual(self.handler.scenarios_path.read_bytes(), original_database)

    def test_table_save_checks_do_not_coerce_cells_or_require_a_complete_table(self):
        for name, table in (
            ('PadRates', {'T01': [0, 0.0, '', None, '001', 'not a number']}),
            ('PadRates', {'T01': [], 'T02': [0]}),
            ('PadRates', {}),
            ('Units', {'volume': 'bbl'}),
            ('DesalinationSurrogate', {'zero': 0, 'float': 1.0, 'text': '001', 'blank': '', 'null': None}),
            ('DesalinationSurrogate', {'INDEX': ['recovery'], 'VALUE': [0.5]}),
        ):
            with self.subTest(table=name, values=table), patch.object(self.handler, 'update_excel', return_value=self.example) as save:
                response = self.client.post('/update_excel', json={
                    'id': 1, 'tableKey': name, 'updatedTable': table, 'revision': None,
                    'clientMetadata': {'enabled': False},
                })
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(save.call_args.args[:2], (1, name))
                self.assertEqual(json.dumps(save.call_args.args[2]), json.dumps(table))

    def test_table_save_zero_id_and_existing_rejections_preserve_storage(self):
        current = self.create_zero_id_scenario()
        table = {'ProductionPads': ['P1'], 'T01': [42], 'T02': [0]}
        for identity, revision in ((0, current['input_revision']), ('0', None), ('0', '')):
            response = self.client.post('/update_excel', json={
                'id': identity, 'tableKey': 'PadRates', 'updatedTable': table, 'revision': revision})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['id'], 0)
            self.assertEqual(response.json(), self.handler.get_scenario(0))
        self.assertEqual(read_inputs(self.handler.get_excelsheet_path(0))['df_parameters']['PadRates'], table)
        workbook = Path(self.handler.get_excelsheet_path(0))
        original_workbook = workbook.read_bytes()
        original_database = self.handler.scenarios_path.read_bytes()
        payload = {'id': 0, 'tableKey': 'PadRates', 'updatedTable': {}}
        for status, overrides in ((404, {'id': 999999}), (400, {'tableKey': 'UnknownTable'}),
                                  (409, {'revision': 'stale-revision'})):
            with self.subTest(status=status):
                response = self.client.post('/update_excel', json={**payload, **overrides})
                self.assertEqual(response.status_code, status, response.text)
        self.handler.add_background_task(0)
        response = self.client.post('/update_excel', json=payload)
        self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(workbook.read_bytes(), original_workbook)
        self.assertEqual(self.handler.scenarios_path.read_bytes(), original_database)

    def test_table_save_response_preserves_legacy_metadata_and_absent_fields(self):
        fixture = Path(__file__).resolve().parents[2] / 'electron/ui/src/tests/fixtures/scenario-contract.json'
        expected = json.loads(fixture.read_text())
        expected['id'] = 1
        with patch.object(self.handler, 'update_excel', return_value=expected):
            response = self.client.post('/update_excel', json={'id': 1, 'tableKey': 'PadRates', 'updatedTable': {}})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(json.dumps(response.json(), sort_keys=True), json.dumps(expected, sort_keys=True))

    def test_table_save_never_returns_malformed_success(self):
        client = TestClient(self.client.app, raise_server_exceptions=False)
        payload = {'id': 1, 'tableKey': 'PadRates', 'updatedTable': {}}
        missing_revision = deepcopy(self.example)
        del missing_revision['input_revision']
        malformed_table = deepcopy(self.example)
        malformed_table['data_input']['df_parameters']['PadRates'] = {'T01': 100}
        invalid_results = deepcopy(self.example)
        invalid_results['results']['data'] = {'flow': 'not rows'}
        for saved in (None, {**self.example, 'id': 0}, {**self.example, 'id': False},
                      missing_revision, {**self.example, 'input_revision': ''},
                      {**self.example, 'input_revision': ' '}, malformed_table, invalid_results):
            with self.subTest(saved=saved), patch.object(self.handler, 'update_excel', return_value=saved):
                self.assertEqual(client.post('/update_excel', json=payload).status_code, 500)

    def test_table_response_failure_can_follow_a_completed_save(self):
        client = TestClient(self.client.app, raise_server_exceptions=False)
        update = self.handler.update_excel
        table = {'ProductionPads': ['P1'], 'T01': [42], 'T02': [0]}

        def lose_revision(*args):
            saved = update(*args)
            del saved['input_revision']
            return saved

        with patch.object(self.handler, 'update_excel', side_effect=lose_revision):
            response = client.post('/update_excel', json={'id': 1, 'tableKey': 'PadRates', 'updatedTable': table})
        self.assertEqual(response.status_code, 500)
        # An acknowledgement failure is not a rollback. The client must reload,
        # retaining its draft in the meantime, rather than automatically retry.
        self.assertEqual(self.handler.get_scenario(1)['data_input']['df_parameters']['PadRates'], table)
        self.assertEqual(read_inputs(self.handler.get_excelsheet_path(1))['df_parameters']['PadRates'], table)

    def test_legacy_units_support_detail_readiness_rename_and_table_save(self):
        path = Path(self.handler.get_excelsheet_path(1))
        legacy_units_sheet(path)
        legacy = deepcopy(self.example)
        expected_units = legacy['data_input'].pop('units')
        legacy['data_input']['df_parameters']['Units'] = deepcopy(expected_units)
        # Simulate persisted v3 data without running today's persistence adapter.
        self.handler._db.update({'scenario': legacy}, Query().id_ == 1)
        self.handler.retrieve_scenarios()
        workbook_before = path.read_bytes()
        database_before = self.handler.scenarios_path.read_bytes()

        response = self.client.get('/get_scenario/1')
        self.assertEqual(response.status_code, 200, response.text)
        current = response.json()
        self.assertEqual(current['data_input']['units'], expected_units)
        response = self.client.get('/scenario_readiness/1')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.json()['valid'], response.text)
        self.assertFalse(any(issue['code'] == 'invalid_units' for issue in response.json()['issues']))
        self.assertEqual(path.read_bytes(), workbook_before)
        self.assertEqual(self.handler.scenarios_path.read_bytes(), database_before)

        current['name'] = 'Renamed legacy scenario'
        response = self.client.post('/update', json={'updatedScenario': current})
        self.assertEqual(response.status_code, 200, response.text)
        current = response.json()['data']
        self.assertEqual(path.read_bytes(), workbook_before)
        table = {'ProductionPads': ['P1'], 'T01': [42], 'T02': [0]}
        response = self.client.post('/update_excel', json={
            'id': 1, 'tableKey': 'PadRates', 'updatedTable': table, 'revision': current['input_revision']})
        self.assertEqual(response.status_code, 200, response.text)
        current = response.json()
        self.assertEqual(current['data_input']['units'], expected_units)
        self.assertEqual(current['data_input']['df_parameters']['Units'], expected_units)
        self.assertEqual(current['data_input']['df_parameters']['PadRates'], table)
        self.assertEqual(current['optimization'], legacy['optimization'])
        self.assertEqual(self.client.get('/get_template/1').content, path.read_bytes())
        workbook = load_workbook(path, read_only=True)
        try:
            self.assertEqual(workbook['Units']['B2'].value, 'Unit')
            self.assertEqual(workbook['Units']['D3'].value, 'Keep this source explanation.')
        finally:
            workbook.close()
        self.assertEqual(read_inputs(path)['units'], expected_units)
        self.assertEqual(self.handler.get_scenario(1), current)

    def test_scalar_metadata_survives_table_save_map_rename_and_model_check(self):
        current = deepcopy(self.example)
        metadata = {'inlet_salinity': 100.0, 'recovery': 0.5, 'source': '001', 'zero': 0}
        current['data_input']['df_parameters']['DesalinationSurrogate'] = metadata
        current['data_input']['df_parameters']['Units'] = deepcopy(current['data_input']['units'])
        current = self.handler.update_scenario(current)
        original_metadata = json.dumps(metadata, sort_keys=True)
        table = {'ProductionPads': ['P1'], 'T01': [42], 'T02': [0]}
        response = self.client.post('/update_excel', json={
            'id': 1, 'tableKey': 'PadRates', 'updatedTable': table, 'revision': current['input_revision']})
        self.assertEqual(response.status_code, 200, response.text)
        current = response.json()
        mapped = current['data_input']['map_data']
        mapped['all_nodes']['Production A'] = mapped['all_nodes'].pop('P1')
        mapped['_node_renames'] = {'P1': 'Production A'}
        mapped['all_nodes']['K1']['DisposalOperationalCost'] = 2
        response = self.client.post('/update', json={'updatedScenario': current, 'propagateChanges': 'map'})
        self.assertEqual(response.status_code, 200, response.text)
        saved = response.json()['data']
        parameters = saved['data_input']['df_parameters']
        self.assertEqual(json.dumps(parameters['DesalinationSurrogate'], sort_keys=True), original_metadata)
        self.assertEqual(parameters['Units'], current['data_input']['df_parameters']['Units'])
        self.assertEqual(parameters['PadRates'], {**table, 'ProductionPads': ['Production A']})
        self.assertEqual(parameters['DisposalOperationalCost']['VALUE'], [2])
        self.assertEqual(saved['optimization'], current['optimization'])
        self.assertEqual(self.handler.get_scenario(1), saved)
        path = Path(self.handler.get_excelsheet_path(1))
        _, model_parameters, _ = get_data(path)
        self.assertEqual(model_parameters['DesalinationSurrogate']['inlet_salinity'], 100)
        self.assertEqual(model_parameters['DesalinationSurrogate']['recovery'], 0.5)
        workbook = load_workbook(path, read_only=True)
        try:
            rows = dict(workbook['DesalinationSurrogate'].iter_rows(min_row=3, max_col=2, values_only=True))
            self.assertEqual(rows, metadata)
        finally:
            workbook.close()
        self.assertEqual(self.client.get('/get_template/1').content, path.read_bytes())
        response = self.client.get('/validate_scenario/1')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['model_check'], 'passed', response.text)

    def test_replacing_a_workbook_does_not_restore_previous_scalar_metadata(self):
        current = deepcopy(self.example)
        current['data_input']['df_parameters']['DesalinationSurrogate'] = {'inlet_salinity': 100, 'recovery': 0.5}
        self.handler.update_scenario(current)
        replacement = Path(self.temp.name) / 'replacement.xlsx'
        write_inputs(self.example['data_input'], replacement)
        with replacement.open('rb') as stream:
            response = self.client.post('/replace/1', files={'file': ('replacement.xlsx', stream)})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotIn('DesalinationSurrogate', response.json()['data_input']['df_parameters'])

    def test_kml_and_shapefile_import_classification_forecasts_and_optimization(self):
        self.handler.update_next_id()
        with patch.dict(os.environ, {'PATH': str(Path.home() / '.idaes/bin') + os.pathsep + os.environ['PATH']}):
            if not SolverFactory('cbc').available(False):
                self.skipTest('CBC is not installed')
            for path in map_files(self.temp.name):
                with self.subTest(format=path.suffix), path.open('rb') as stream:
                    response = self.client.post('/upload/Imported%20network?defaultNodeType=NetworkNode', files={'file': (path.name, stream)})
                    self.assertEqual(response.status_code, 200, response.text)
                    scenario = response.json()
                    scenario_id = scenario['id']
                    mapped = scenario['data_input']['map_data']
                    self.assertEqual(set(mapped['all_nodes']), {'P1', 'N1', 'K1'})
                    for name, node in self.example['data_input']['map_data']['all_nodes'].items():
                        mapped['all_nodes'][name].update(node)
                    for arc in mapped['arcs'].values():
                        arc['diameter'] = 'D4'
                        for node in arc['nodes']:
                            node['outgoing_nodes'] = {'P1': ['N1'], 'N1': ['K1'], 'K1': []}[node['name']]
                    response = self.client.post('/update', json={'updatedScenario': scenario, 'propagateChanges': 'map'})
                    self.assertEqual(response.status_code, 200, response.text)
                    response = self.client.post(f'/planning_horizon/{scenario_id}', json={'periods': ['T01', 'T02']})
                    self.assertEqual(response.status_code, 200, response.text)
                    response = self.client.post('/update_excel', json={'id': scenario_id, 'tableKey': 'PadRates',
                        'updatedTable': {'ProductionPads': ['P1'], 'T01': [100], 'T02': [0]}})
                    self.assertEqual(response.status_code, 200, response.text)
                    scenario = response.json()
                    validation = self.client.get(f'/validate_scenario/{scenario_id}').json()
                    self.assertTrue(validation['valid'], validation)
                    scenario['optimization']['runtime'] = 20
                    response = self.client.post('/run_model', json={'scenario': scenario})
                    self.assertEqual(response.status_code, 200, response.text)
                    completed = self.handler.get_scenario(scenario_id)
                    self.assertEqual(completed['results']['status'], 'Optimized', completed['results'].get('error'))
                    self.assertEqual(completed['results']['solution_status'], 'optimal')
                    self.assertEqual(completed['data_input']['df_parameters']['PadRates']['T02'], [0])

    def test_complete_optimization_produces_report_from_validated_inputs(self):
        with patch.dict(os.environ, {'PATH': str(Path.home() / '.idaes/bin') + os.pathsep + os.environ['PATH']}):
            if not SolverFactory('cbc').available(False):
                self.skipTest('CBC is not installed')
            validation = self.client.post('/scenario_feasibility/1')
            self.assertEqual(validation.json()['feasibility'], 'feasible', validation.text)
            self.assertEqual(self.client.post('/advance_to_optimization_setup/1').status_code, 200)
            current = self.handler.get_scenario(1)
            current['optimization']['runtime'] = 20
            with patch.object(self.runner, 'create_model', wraps=self.runner.create_model) as build:
                response = self.client.post('/run_model', json={'scenario': current, 'run_id': 'real-solve'})
            build.assert_called_once()
            self.assertEqual(response.status_code, 200, response.text)
        result = self.handler.get_scenario(1)
        self.assertEqual(result['results']['status'], 'Optimized', result['results'].get('error'))
        self.assertEqual(result['results']['solution_status'], 'optimal')
        self.assertEqual(result['results']['input_revision'], result['input_revision'])
        self.assertEqual(result['results']['run_id'], 'real-solve')
        self.assertGreater(len(result['results']['data']), 50)
        report = self.client.get('/generate_report/1')
        self.assertEqual(report.status_code, 200)
        self.assertTrue(report.content.startswith(b'PK'))
        self.assertEqual(self.handler.get_background_tasks(), [])
        self.assertFalse(list(self.handler.excelsheets_path.glob('optimization-*')))

    def test_run_endpoint_rechecks_changed_settings(self):
        self.client.get('/validate_scenario/1')
        current = self.handler.get_scenario(1)
        current['results'] = {'status': 'Optimized', 'data': {'previous': [[123]]}}
        self.handler.update_scenario(current)
        current['optimization']['waterQuality'] = 'discrete'
        response = self.client.post('/run_model', json={'scenario': current})
        self.assertEqual(response.status_code, 422, response.text)
        self.assertTrue(any(issue['table'] == 'PadWaterQuality' for issue in response.json()['detail']['validation']['issues']))
        self.assertEqual(self.handler.get_background_tasks(), [])
        self.assertEqual(self.handler.get_scenario(1)['results'], current['results'])

    def test_launch_acknowledges_before_preparation_and_reserves_an_immutable_snapshot(self):
        supplied = deepcopy(self.example)
        async def payload():
            return {'scenario': supplied, 'run_id': 'launch-1'}
        tasks = BackgroundTasks()
        with patch.object(self.runner, 'write_inputs') as write, patch.object(self.runner, 'create_model') as build, \
                patch.object(self.handler, 'validate__pareto_scenario') as validate_model:
            accepted = asyncio.run(self.routes.run_model(SimpleNamespace(json=payload), tasks))
            write.assert_not_called()
            build.assert_not_called()
            validate_model.assert_not_called()
        self.assertEqual(accepted['results']['status'], 'Preparing inputs')
        self.assertEqual(self.handler.get_background_tasks(), [1])
        self.assertEqual(len(tasks.tasks), 1)
        captured = deepcopy(tasks.tasks[0].kwargs)
        supplied['data_input']['df_parameters']['PadRates']['T01'] = [999]
        supplied['optimization']['runtime'] = 999
        supplied['override_values'] = {'changed': {}}
        accepted['data_input']['df_parameters']['PadRates']['T02'] = [999]
        self.assertEqual(tasks.tasks[0].kwargs, captured)

        # Retry a lost acknowledgement with the same identity, without a second task.
        repeated_tasks = BackgroundTasks()
        repeated = asyncio.run(self.routes.run_model(SimpleNamespace(json=payload), repeated_tasks))
        self.assertEqual(repeated['results']['run_id'], 'launch-1')
        self.assertFalse(repeated_tasks.tasks)
        async def duplicate():
            return {'scenario': self.example, 'run_id': 'different-launch'}
        with self.assertRaises(HTTPException) as error:
            asyncio.run(self.routes.run_model(SimpleNamespace(json=duplicate), BackgroundTasks()))
        self.assertEqual(error.exception.status_code, 409)

    def test_preparation_runs_outside_the_event_loop_and_database_lock(self):
        entered, release = threading.Event(), threading.Event()
        def slow_write(*args):
            entered.set()
            if not release.wait(5):
                raise RuntimeError('Test preparation timed out')
            raise RuntimeError('Stop before solving')
        async def exercise():
            async def payload():
                return {'scenario': self.example}
            tasks = BackgroundTasks()
            await self.routes.run_model(SimpleNamespace(json=payload), tasks)
            worker = asyncio.create_task(tasks())
            try:
                self.assertTrue(await asyncio.to_thread(entered.wait, 5))
                current = await asyncio.wait_for(self.routes.get_scenario('1'), 1)
                active = await asyncio.wait_for(self.routes.check_tasks(), 1)
                self.assertEqual(current['results']['status'], 'Preparing inputs')
                self.assertEqual(active, {'tasks': [1]})
                self.assertFalse(release.is_set())
            finally:
                release.set()
                await worker
        with patch.object(self.runner, 'write_inputs', side_effect=slow_write):
            asyncio.run(exercise())
        self.assertEqual(self.handler.get_background_tasks(), [])

    def test_preparation_failures_record_the_stage_and_release_snapshot_and_reservation(self):
        for function, stage in [('write_inputs', 'Preparing inputs'), ('create_model', 'Building model')]:
            with self.subTest(stage=stage), patch.object(self.runner, function, side_effect=RuntimeError('Cannot prepare these inputs')):
                current = self.handler.get_scenario(1)
                response = self.client.post('/run_model', json={'scenario': current, 'run_id': function})
            self.assertEqual(response.status_code, 200, response.text)
            failed = self.handler.get_scenario(1)
            self.assertEqual(failed['results']['status'], 'failure')
            self.assertEqual(failed['results']['failure_stage'], stage)
            self.assertEqual(failed['results']['run_id'], function)
            self.assertEqual(failed['results']['error'], 'Cannot prepare these inputs')
            self.assertEqual(self.handler.get_background_tasks(), [])
            self.assertFalse(list(self.handler.excelsheets_path.glob('optimization-*')))
            if stage == 'Building model':
                self.assertEqual(failed['validation']['model_check'], 'failed')

    def test_explicit_facility_rename_preserves_forecasts_and_connections(self):
        scenario = deepcopy(self.example)
        mapped = scenario['data_input']['map_data']
        mapped['all_nodes']['Production A'] = mapped['all_nodes'].pop('P1')
        mapped['_node_renames'] = {'P1': 'Production A'}
        response = self.client.post('/update', json={'updatedScenario': scenario, 'propagateChanges': 'map'})
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()['data']['data_input']
        self.assertEqual(data['df_parameters']['PadRates']['ProductionPads'], ['Production A'])
        self.assertEqual(data['df_parameters']['PadRates']['T02'], [100])
        self.assertEqual(data['df_parameters']['PNA']['N1'], [1])
        self.assertNotIn('P1', data['map_data']['all_nodes'])

    def test_validation_finishing_after_an_edit_is_not_saved_as_current(self):
        current = deepcopy(self.example)
        self.handler.update_horizon(1, ['T01'])
        result = self.handler._save_validation_results(current, {'valid': True, 'revision': current['input_revision']})
        self.assertEqual(result['state'], 'outdated')
        self.assertEqual(self.handler.get_scenario(1)['validation']['state'], 'inputs_complete')
        self.assertEqual(self.handler.get_scenario(1)['validation']['model_check'], 'not_run')

    def test_saving_one_table_keeps_other_errors_and_partial_fixes_highlighted(self):
        scenario = deepcopy(self.example)
        tables = scenario['data_input']['df_parameters']
        tables['PadRates']['T01'] = ['']
        tables['PadRates']['T02'] = ['']
        tables['DisposalOperationalCost']['VALUE'] = ['']
        self.handler.save_inputs(scenario)
        self.client.get('/validate_scenario/1')
        tables['PadRates']['T01'] = [100]
        response = self.client.post('/update_excel', json={'id': 1, 'tableKey': 'PadRates', 'updatedTable': tables['PadRates']})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn('PadRates', response.json()['validation']['tables_with_issues'])
        tables['PadRates']['T02'] = [100]
        response = self.client.post('/update_excel', json={'id': 1, 'tableKey': 'PadRates', 'updatedTable': tables['PadRates']})
        checked = response.json()['validation']
        self.assertNotIn('PadRates', checked['tables_with_issues'])
        self.assertIn('DisposalOperationalCost', checked['tables_with_issues'])
        self.assertEqual(self.handler.get_scenario(1)['validation'], checked)

    def test_valid_edit_refreshes_inputs_but_requires_a_new_model_check(self):
        self.client.get('/validate_scenario/1')
        table = {'ProductionPads': ['P1'], 'T01': [90], 'T02': [100]}
        response = self.client.post('/update_excel', json={'id': 1, 'tableKey': 'PadRates', 'updatedTable': table})
        result = response.json()['validation']
        self.assertTrue(result['valid'])
        self.assertEqual(result['state'], 'inputs_complete')
        self.assertEqual(result['model_check'], 'not_run')
        self.assertEqual(result['feasibility'], 'not_run')
        self.assertEqual(self.client.post('/advance_to_optimization_setup/1').status_code, 409)

    def test_autofill_preview_apply_export_and_concurrent_edit_protection(self):
        scenario = deepcopy(self.example)
        scenario['data_input']['df_parameters']['DisposalOperationalCost']['VALUE'] = ['']
        current = self.handler.save_inputs(scenario)
        payload = {'section': 'costs', 'value': 2, 'revision': current['input_revision']}
        preview = self.client.post('/fill_scenario_inputs/1', json=payload)
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertGreater(preview.json()['cell_count'], 0)
        self.assertEqual(self.handler.get_scenario(1)['input_revision'], current['input_revision'])
        saved = self.client.post('/fill_scenario_inputs/1', json={**payload, 'apply': True})
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertNotIn('DisposalOperationalCost', saved.json()['validation']['tables_with_issues'])
        self.assertEqual(saved.json()['data_input']['map_data']['all_nodes']['K1']['DisposalOperationalCost'], 2)
        self.assertEqual(read_inputs(self.handler.get_excelsheet_path(1))['df_parameters']['DisposalOperationalCost']['VALUE'], [2])
        self.assertEqual(self.client.post('/fill_scenario_inputs/1', json={**payload, 'apply': True}).status_code, 409)
        self.handler.add_background_task(1)
        self.assertEqual(self.client.post('/fill_scenario_inputs/1', json={**payload, 'revision': saved.json()['input_revision'], 'apply': True}).status_code, 409)

    def test_autofill_invalid_value_does_not_partially_save(self):
        original = self.handler.get_scenario(1)
        response = self.client.post('/fill_scenario_inputs/1', json={
            'section': 'capacity', 'revision': original['input_revision'], 'value': -2, 'apply': True})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.handler.get_scenario(1), original)

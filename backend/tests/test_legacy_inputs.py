"""Saved v3 metadata must survive workbook and map adapters without migration."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from app.internal.scenarios.inputs import read_inputs, write_inputs, prune_removed_map_nodes, rename_map_nodes
from app.internal.scenarios.input_schema import input_revision
from app.internal.workbooks.reader import get_data
from app.internal.validation.scenario_validation import validate_inputs
from scenario_fixtures import legacy_units_sheet, simple_scenario


class LegacyInputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.scenario, self.path = simple_scenario(self.temp.name)
        # The fixture's last write clears index-only optional tables. Start from
        # that saved representation before comparing revisions across our edits.
        self.scenario['data_input'] = read_inputs(self.path, previous=self.scenario['data_input'])

    def test_legacy_unit_labels_load_without_changing_the_workbook(self):
        legacy_units_sheet(self.path)
        original = self.path.read_bytes()
        expected = self.scenario['data_input']['units']
        _, parameters, _ = get_data(self.path)
        self.assertEqual(parameters['Units'], expected)
        restored = read_inputs(self.path)
        self.assertEqual(restored['units'], expected)
        self.assertEqual(input_revision({**self.scenario, 'data_input': restored}),
                         input_revision(self.scenario))
        self.assertEqual(self.path.read_bytes(), original)

    def test_missing_legacy_unit_stays_blank_instead_of_getting_a_default(self):
        legacy_units_sheet(self.path)
        workbook = load_workbook(self.path)
        try:
            workbook['Units']['B3'] = None
            workbook.save(self.path)
        finally:
            workbook.close()
        _, parameters, _ = get_data(self.path)
        self.assertEqual(parameters['Units']['volume'], '')
        # JSON hashing must work so ordinary readiness checks can report the blank.
        input_revision({'data_input': {'units': parameters['Units']}})
        self.scenario['data_input']['units'] = parameters['Units']
        validation = validate_inputs(self.scenario)
        self.assertTrue(any(issue['code'] == 'invalid_units' and issue['row'] == ['volume']
                            for issue in validation['issues']))

    def test_scalar_metadata_survives_saved_and_fresh_workbook_round_trips(self):
        data = self.scenario['data_input']
        metadata = {'inlet_salinity': 100.0, 'recovery': 0.5, 'code': '001',
                    'numeric_text': '1.0', 'zero': 0, 'blank': '', 'missing': None}
        data['df_parameters']['DesalinationSurrogate'] = metadata
        data['df_parameters']['Units'] = deepcopy(data['units'])
        original = json.dumps(data, sort_keys=True)
        for path in (self.path, Path(self.temp.name) / 'run-snapshot.xlsx'):
            with self.subTest(existing=path == self.path):
                write_inputs(data, path)
                restored = read_inputs(path, previous=data)
                self.assertEqual(json.dumps(restored['df_parameters']['DesalinationSurrogate'], sort_keys=True),
                                 json.dumps(metadata, sort_keys=True))
                self.assertEqual(restored['df_parameters']['Units'], data['df_parameters']['Units'])
                self.assertEqual(input_revision({'data_input': restored}), input_revision({'data_input': data}))
                _, parameters, _ = get_data(path)
                self.assertEqual(parameters['DesalinationSurrogate']['inlet_salinity'], 100)
                self.assertEqual(parameters['DesalinationSurrogate']['recovery'], 0.5)
                self.assertEqual(parameters['DesalinationSurrogate']['zero'], 0)
                workbook = load_workbook(path, read_only=True)
                try:
                    rows = dict(workbook['DesalinationSurrogate'].iter_rows(min_row=3, max_col=2, values_only=True))
                    self.assertEqual(rows, {**metadata, 'blank': None})
                finally:
                    workbook.close()
        self.assertEqual(json.dumps(data, sort_keys=True), original)

    def test_map_pruning_and_renames_leave_scalar_metadata_intact(self):
        data = self.scenario['data_input']
        metadata = {'inlet_salinity': 100, 'recovery': 0.5, 'INDEX': 'P1'}
        data['df_parameters']['DesalinationSurrogate'] = metadata
        data['df_parameters']['Units'] = deepcopy(data['units'])
        original = deepcopy(data)
        mapped = deepcopy(data['map_data'])
        del mapped['all_nodes']['P1']
        pruned = prune_removed_map_nodes(data, mapped)
        self.assertEqual(pruned['df_parameters']['PadRates']['ProductionPads'], [])
        renamed = rename_map_nodes(data, {'P1': 'P2'})
        self.assertEqual(renamed['df_parameters']['PadRates']['ProductionPads'], ['P2'])
        for result in (pruned, renamed):
            self.assertEqual(result['df_parameters']['DesalinationSurrogate'], metadata)
            self.assertEqual(result['df_parameters']['Units'], data['df_parameters']['Units'])
        self.assertEqual(data, original)

    def test_metadata_support_does_not_accept_malformed_ordinary_tables(self):
        original = self.path.read_bytes()
        for name, table in (
            ('PadRates', {'ProductionPads': ['P1'], 'T01': 100}),
            ('DesalinationSurrogate', {'inlet_salinity': [100], 'recovery': 0.5}),
        ):
            with self.subTest(name=name):
                data = deepcopy(self.scenario['data_input'])
                data['df_parameters'][name] = table
                with self.assertRaises(ValueError):
                    write_inputs(data, self.path)
                self.assertEqual(self.path.read_bytes(), original)
                validation = validate_inputs({**self.scenario, 'data_input': data})
                self.assertTrue(any(issue['code'] == 'invalid_table' and issue['table'] == name
                                    for issue in validation['issues']))


if __name__ == '__main__':
    unittest.main()

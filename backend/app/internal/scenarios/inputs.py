"""Canonical saved inputs and immutable workbooks for model runs."""
from copy import deepcopy
from pathlib import Path
import os
import tempfile
from openpyxl import load_workbook
from pareto.utilities.get_data import get_display_units
from app.internal.workbooks.reader import get_data, get_input_lists
from app.internal.scenarios.input_schema import DEFAULT_UNITS, FORECASTS, NODE_FIELDS, NODE_SETS, flat_table, dimension_count, is_scalar_parameter

def read_inputs(path, previous=None, map_data=None):
    sets, parameters, tables = get_data(str(path), *get_input_lists())
    data = deepcopy(previous or {})
    # A save/map rebuild writes these values as INDEX/VALUE rows for PARETO.
    # Keep their saved JSON representation, including numeric strings and nulls,
    # when reading that workbook back. A fresh import has no previous metadata.
    metadata = {name: table for name, table in data.get('df_parameters', {}).items()
                if is_scalar_parameter(name, table)}
    data.update(df_sets={k: list(v) for k, v in sets.items()},
                df_parameters={**{k: v for k, v in tables.items() if k != 'Units'}, **metadata},
                units=parameters['Units'],
                display_units=get_display_units(get_input_lists()[1], parameters['Units']))
    if map_data is not None:
        data['map_data'] = deepcopy(map_data)
    sync_map_fields(data)
    return data

def sync_map_fields(data):
    """Tables own saved numeric values; keep map forms in sync without losing geometry."""
    map_data = data.get('map_data')
    if not map_data:
        return
    types = {node: kind for key, kind in NODE_SETS.items() for node in data.get('df_sets', {}).get(key, [])}
    fields = {key: flat_table(data.get('df_parameters', {}).get(key, {})) for key in NODE_FIELDS}
    for name, node in map_data.get('all_nodes', {}).items():
        if name in types:
            node['node_type'] = types[name]
            node['nodeType'] = types[name]
        for key, values in fields.items():
            if (name,) in values:
                node[key] = values[(name,)]
    map_data['units'] = {k.replace(' ', '_'): v for k, v in data.get('units', DEFAULT_UNITS).items()}
    from pareto.utilities.process_data import get_valid_piping_arc_list
    edges = {key for table in get_valid_piping_arc_list()
             for key, value in flat_table(data.get('df_parameters', {}).get(table, {})).items()
             if value not in ('', None, 0, '0', False)}
    for arc in map_data.get('arcs', {}).values():
        nodes = arc.get('nodes', [])
        for i, node in enumerate(nodes):
            adjacent = nodes[max(0, i - 1):i] + nodes[i + 1:i + 2]
            node['outgoing_nodes'] = [other['name'] for other in adjacent if (node['name'], other['name']) in edges]

def prune_removed_map_nodes(data, map_data):
    """Explicit map removal also removes references to those facilities in tables."""
    result = deepcopy(data)
    nodes = map_data.get('all_nodes', {})
    allowed = set(nodes)
    from pareto.utilities.process_data import get_valid_piping_arc_list, get_valid_trucking_arc_list
    node_columns = set(get_valid_piping_arc_list() + get_valid_trucking_arc_list()) | {
        'InitialPipelineCapacity', 'InitialPipelineDiameters', 'PipelineOperationalCost', 'PipelineExpansionDistance', 'TruckingTime'}
    new_sets = {key: {name for name, node in nodes.items() if node.get('node_type') == kind} for key, kind in NODE_SETS.items()}
    for name, table in result.get('df_parameters', {}).items():
        if not table or is_scalar_parameter(name, table):
            continue
        headers = list(table)
        indices = headers[:dimension_count(headers)]
        keep = []
        for row in range(len(table[headers[0]])):
            valid = True
            for header in indices:
                base = header.split('.')[0]
                domain = new_sets.get(base)
                if base.lower() in ('nodes', 'pads'):
                    domain = allowed
                if domain is not None and table[header][row] not in domain:
                    valid = False
            if valid:
                keep.append(row)
        result['df_parameters'][name] = {header: [values[i] for i in keep if i < len(values)]
            for header, values in table.items()
            if name not in node_columns or header in indices or header in allowed}
    return result

def rename_map_nodes(data, renames):
    """Apply explicit identifier changes to geometry and indexed table values."""
    result = deepcopy(data)
    for key in NODE_SETS:
        result.get('df_sets', {})[key] = [renames.get(n, n) for n in result.get('df_sets', {}).get(key, [])]
    from pareto.utilities.process_data import get_valid_piping_arc_list, get_valid_trucking_arc_list
    matrices = set(get_valid_piping_arc_list() + get_valid_trucking_arc_list()) | {
        'InitialPipelineCapacity', 'InitialPipelineDiameters', 'PipelineOperationalCost', 'PipelineExpansionDistance', 'TruckingTime'}
    for name, table in result.get('df_parameters', {}).items():
        if is_scalar_parameter(name, table):
            continue
        headers = list(table)
        indices = headers[:dimension_count(headers)]
        result['df_parameters'][name] = {
            renames.get(header, header) if name in matrices and header not in indices else header:
            [renames.get(value, value) for value in values] if header in indices and (header.split('.')[0] in NODE_SETS or header.lower().split('.')[0] in ('nodes', 'pads')) else values
            for header, values in table.items()}
    mapped = result.get('map_data', {})
    for key in ('all_nodes', *NODE_SETS):
        if key in mapped:
            mapped[key] = {renames.get(n, n): value for n, value in mapped[key].items()}
    for arc in mapped.get('arcs', {}).values():
        for node in arc.get('nodes', []):
            node['name'] = renames.get(node['name'], node['name'])
            node['outgoing_nodes'] = [renames.get(n, n) for n in node.get('outgoing_nodes', [])]
    return result

def write_inputs(data, path, template=None):
    """Write exactly these tables, sets and units; replace the workbook atomically."""
    path = Path(path)
    from app.internal.workbooks.excel_api import DEFAULT_TEMPLATE_LOCATION
    source = template or (path if path.exists() else DEFAULT_TEMPLATE_LOCATION)
    wb = load_workbook(source, data_only=True)
    try:
        entries = {**data.get('df_sets', {}), **data.get('df_parameters', {})}
        # The legacy parent reader also needs these headers for a CP-free network.
        periods = data.get('df_sets', {}).get('TimePeriods', [])
        demand = entries.get('CompletionsDemand', {})
        if not demand and periods:
            entries['CompletionsDemand'] = {'CompletionsPads': [], **{t: [] for t in periods}}
        for name, table in entries.items():
            if name in ('Units', 'proprietary_data'):
                continue
            if is_scalar_parameter(name, table):
                # Serialize metadata even for a new validation/run workbook;
                # skipping it would silently substitute the model's defaults.
                table = {'INDEX': list(table), 'VALUE': list(table.values())}
            ws = wb[name] if name in wb else wb.create_sheet(name)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.value = None
            if isinstance(table, list):
                ws.cell(1, 1).value = name
                for row, value in enumerate(table, 2):
                    ws.cell(row, 1).value = value
            elif isinstance(table, dict):
                for col, (header, values) in enumerate(table.items(), 1):
                    if not isinstance(values, list):
                        raise ValueError(f'{name}: expected a list of values for {header}.')
                    ws.cell(2, col).value = header
                    for row, value in enumerate(values, 3):
                        ws.cell(row, col).value = None if value == '' else value
            else:
                raise ValueError(f'{name}: invalid table format.')
        units = data.get('units')
        if units:
            ws = wb['Units']
            for row in range(3, ws.max_row + 1):
                if ws.cell(row, 1).value in units:
                    ws.cell(row, 2).value = units[ws.cell(row, 1).value]
        path.parent.mkdir(parents=True, exist_ok=True)
        fd, temp = tempfile.mkstemp(suffix='.xlsx', dir=path.parent)
        os.close(fd)
        try:
            wb.save(temp)
            wb.close()
            os.replace(temp, path)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)
    finally:
        wb.close()

def changed_pipe_fields(current, previous):
    """A diameter edit changes capacity; unrelated map edits do not."""
    def metadata(data):
        result = {}
        for arc in (data or {}).get('arcs', {}).values():
            nodes = arc.get('nodes', [])
            lengths = arc.get('lengths', [])
            for i, (a, b) in enumerate(zip(nodes, nodes[1:])):
                for source, dest in ((a, b), (b, a)):
                    if dest['name'] in source.get('outgoing_nodes', []):
                        result[(source['name'], dest['name'])] = (arc.get('diameter'), lengths[i] if i < len(lengths) else None)
        return result
    old, new = metadata(previous), metadata(current)
    changes = set()
    for edge, (diameter, length) in new.items():
        before = old.get(edge, (None, None))
        if diameter != before[0]:
            changes.update((edge, key) for key in ('pipeline_capacity', 'pipeline_diameter'))
        if length != before[1]:
            changes.add((edge, 'pipeline_length'))
    return changes

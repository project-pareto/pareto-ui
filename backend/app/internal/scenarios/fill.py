"""Preview and apply one explicit value to currently flagged numeric inputs."""
from collections import Counter
from copy import deepcopy
from functools import lru_cache
import math

from app.internal.scenarios.input_schema import FORECASTS, dimension_count, input_revision
from app.internal.validation.scenario_validation import REQUIREMENTS, SECTION_NAMES, numeric, validate_inputs


@lru_cache(maxsize=1)
def template_headers():
    from openpyxl import load_workbook
    from app.internal.workbooks.excel_api import DEFAULT_TEMPLATE_LOCATION
    workbook = load_workbook(DEFAULT_TEMPLATE_LOCATION, read_only=True, data_only=True)
    try:
        result = {}
        for sheet in workbook:
            headers = [str(cell.value) for cell in next(sheet.iter_rows(min_row=2, max_row=2)) if cell.value is not None]
            seen = Counter()
            result[sheet.title] = []
            for header in headers:
                result[sheet.title].append(f'{header}.{seen[header]}' if seen[header] else header)
                seen[header] += 1
        return result
    finally:
        workbook.close()


def set_cell(tables, table, keys, value, periods):
    """Match scalar and matrix indices; create missing rows without guessing names."""
    columns = tables.get(table)
    if not columns:
        headers = ([FORECASTS[table][0], *periods] if table in FORECASTS else template_headers().get(table, []))
        # Older bundled workbooks omit some supported scalar inputs, including
        # InitialStorageLevel. The requirement defines their index explicitly.
        if not headers:
            requirement = next((item for item in REQUIREMENTS if item.table == table), None)
            if requirement:
                headers = [requirement.node_set, 'VALUE']
        # Empty matrix templates contain only their index headings.
        if headers and headers[-1].lower() != 'value' and len(headers) < len(keys):
            headers = [*headers, keys[-1]]
        columns = {header: [] for header in headers}
    if (not isinstance(columns, dict) or len(columns) < 2 or
            any(not isinstance(v, list) for v in columns.values()) or
            len({len(v) for v in columns.values()}) != 1):
        raise ValueError(f'{table}: repair the table columns before filling values.')
    headers = list(columns)
    count = dimension_count(headers)
    scalar = len(headers) == count + 1 and headers[-1].lower() == 'value'
    if len(keys) != count + (0 if scalar else 1):
        raise ValueError(f'{table}: repair the table headings before filling values.')
    column = headers[-1] if scalar else keys[-1]
    if column in headers[:count]:
        raise ValueError(f'{table}: a value column conflicts with an index heading.')
    size = len(columns[headers[0]])
    matches = [i for i in range(size) if tuple(columns[h][i] for h in headers[:count]) == keys[:count]]
    if len(matches) > 1:
        raise ValueError(f'{table}: remove duplicate rows for {" / ".join(keys[:count])} before filling values.')
    if column not in columns:
        columns[column] = [''] * size
    index = matches[0] if matches else size
    if not matches:
        for h in columns:
            columns[h].append(keys[headers.index(h)] if h in headers[:count] else '')
    columns[column][index] = value
    tables[table] = columns


def prepare_fill(scenario, section, value):
    """Compute preview and apply from the same complete set of eligible cells.

    Validation caps visible issues at 250, but collects fill targets separately.
    The route checks the revision before recomputing this operation, so accepting
    a preview cannot apply a value to inputs changed since that preview.
    """
    if section not in SECTION_NAMES:
        raise ValueError('Choose a completion section to fill.')
    number = numeric(value)
    if number is None or isinstance(value, bool):
        raise ValueError('Enter a finite numeric value.')
    targets = {}
    result = validate_inputs(scenario, fill_targets=targets)
    selected = [t for t in targets.values() if t['section'] == section]
    for target in selected:
        if (not target['minimum'] <= number <= target['maximum'] or
                target['table'] == 'ReuseCapacity' and -1 < number < 0):
            bound = (f"at least {target['minimum']}" if target['maximum'] == math.inf else
                     f"between {target['minimum']} and {target['maximum']}")
            if target['table'] == 'ReuseCapacity':
                bound = 'nonnegative, or -1 for unrestricted capacity'
            raise ValueError(f"{target['table']} requires a value {bound}. Fill tables individually if they need different values.")
    updated = deepcopy(scenario)
    tables = updated['data_input'].setdefault('df_parameters', {})
    for target in selected:
        set_cell(tables, target['table'], target['keys'], number, result['periods'])
    counts = Counter(t['table'] for t in selected)
    units = scenario['data_input'].get('display_units', {})
    return updated, {
        'revision': input_revision(scenario), 'value': number, 'cell_count': len(selected),
        'tables': [{'name': name, 'cell_count': count, 'unit': units.get(name, '')} for name, count in sorted(counts.items())],
    }
